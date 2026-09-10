#!/usr/bin/env python3
"""Extract FM2 Ghost v1.3.4 duelists and drop pools from the data-dump workbook.

Source workbook sheets used:
  - Cards List
  - Drop Pools
  - Drop Pool Unlock Requirements

Outputs:
  - duelists.json
  - drops.json
  - drop-unlock-requirements.json

Notes:
  * Drop weights are preserved exactly as the workbook stores them (often fractional)
    and use denominator=2048.
  * Pink-highlighted card names are marked as requiring at least 400 wins.
  * Yellow-highlighted card names are marked as requiring the card to already be
    in the player's library.
  * Cyan-highlighted Kuriboh SA POW entries are marked as participating in the
    workbook's Kuriboh drop-progression mechanic.
  * Unlock columns are alternate versions of one rank pool. To keep drops.json
    compact, unchanged entries are emitted once. Added/changed unlock entries are
    emitted with the unlock condition, while base-only/changed entries are marked
    as default-pool variants.
"""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

DROP_DENOMINATOR = 2048

# Google Sheets colors as exported to XLSX.
PINK_FILL = "FFC27BA0"     # requires 400 wins
YELLOW_FILL = "FFFFFF00"   # requires already owning card in library
CYAN_FILL = "FF00FFFF"     # Kuriboh POW questline marker

RANK_MAP = {
    "SA POW": "SA_POW",
    "BCD": "BCD",
    "SA TEC": "SA_TEC",
}

# Two spelling inconsistencies in the Drop Pools sheet versus Cards List.
CARD_NAME_ALIASES = {
    "Botanical Leon": "Botanical Lion",
    "Giga Plant": "Gigaplant",
}


def slugify(text: str) -> str:
    text = unicodedata.normalize("NFKD", text)
    text = text.encode("ascii", "ignore").decode("ascii")
    text = text.lower().replace("'", "")
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return text


def fill_rgb(cell) -> str | None:
    if cell.fill.fill_type != "solid":
        return None
    color = cell.fill.fgColor
    if color.type == "rgb":
        return color.rgb.upper() if color.rgb else None
    return None


def cell_special_conditions(cell, rank: str) -> list[str]:
    fill = fill_rgb(cell)
    conditions: list[str] = []
    if fill == PINK_FILL:
        conditions.append("Requires at least 400 wins")
    elif fill == YELLOW_FILL:
        conditions.append("Requires this card to already be in the library")
    elif fill == CYAN_FILL and rank == "SA_POW":
        conditions.append("Kuriboh POW questline may transform this drop based on trunk counts")
    return conditions


def combine_conditions(*parts: str | None) -> str | None:
    cleaned = [p.strip() for p in parts if p and p.strip()]
    return "; ".join(cleaned) if cleaned else None


def load_card_lookup(ws) -> tuple[dict[str, int], dict[int, str]]:
    name_to_id: dict[str, int] = {}
    id_to_name: dict[int, str] = {}
    for row in range(2, ws.max_row + 1):
        raw_id = ws.cell(row, 1).value
        raw_name = ws.cell(row, 2).value
        if raw_id is None or raw_name is None:
            continue
        card_id = int(raw_id)
        name = str(raw_name).strip()
        if name in name_to_id:
            raise ValueError(f"Duplicate card name in Cards List: {name!r}")
        if card_id in id_to_name:
            raise ValueError(f"Duplicate card ID in Cards List: {card_id}")
        name_to_id[name] = card_id
        id_to_name[card_id] = name
    return name_to_id, id_to_name


def resolve_card_id(name: str, name_to_id: dict[str, int]) -> tuple[int, str]:
    source_name = name.strip()
    canonical = CARD_NAME_ALIASES.get(source_name, source_name)
    if canonical not in name_to_id:
        raise KeyError(f"Drop card name not found in Cards List: {source_name!r}")
    return name_to_id[canonical], canonical


def parse_unlock_requirements(ws) -> dict[str, list[str]]:
    requirements: dict[str, list[str]] = {}
    row = 1
    while row <= ws.max_row:
        value = ws.cell(row, 1).value
        if value is None:
            row += 1
            continue
        duelist = str(value).strip()
        reqs: list[str] = []
        row += 1
        while row <= ws.max_row:
            value = ws.cell(row, 1).value
            if value is None:
                break
            reqs.append(str(value).strip())
            row += 1
        requirements[duelist] = reqs
        row += 1
    return requirements


def find_duelist_sections(ws) -> list[tuple[int, str]]:
    sections: list[tuple[int, str]] = []
    for row in range(1, ws.max_row):
        name = ws.cell(row, 1).value
        next_value = ws.cell(row + 1, 1).value
        if isinstance(name, str) and next_value == "SA POW":
            sections.append((row, name.strip()))
    return sections


def read_pool(ws, start_row: int, end_row: int, name_col: int, rank: str,
              name_to_id: dict[str, int]) -> dict[int, dict[str, Any]]:
    pool: dict[int, dict[str, Any]] = {}
    for row in range(start_row, end_row + 1):
        name_value = ws.cell(row, name_col).value
        weight_value = ws.cell(row, name_col + 1).value
        if name_value is None:
            continue
        if weight_value is None:
            raise ValueError(f"Missing weight for {name_value!r} at row {row}, col {name_col + 1}")
        source_name = str(name_value).strip()
        card_id, canonical_name = resolve_card_id(source_name, name_to_id)
        if card_id in pool:
            raise ValueError(f"Duplicate card {canonical_name!r} inside one drop pool")
        special = cell_special_conditions(ws.cell(row, name_col), rank)
        pool[card_id] = {
            "cardId": card_id,
            "sourceName": source_name,
            "canonicalName": canonical_name,
            "weight": float(weight_value),
            "specialCondition": "; ".join(special) if special else None,
        }
    return pool


def make_drop(duelist_id: int, rank: str, item: dict[str, Any], condition: str | None = None) -> dict[str, Any]:
    out: dict[str, Any] = {
        "duelistId": duelist_id,
        "cardId": item["cardId"],
        "rank": rank,
        "weight": item["weight"],
        "denominator": DROP_DENOMINATOR,
    }
    combined = combine_conditions(item.get("specialCondition"), condition)
    if combined:
        out["condition"] = combined
    return out


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", nargs="?", default="Forbidden Memories 2 v1.3.4 data dump.xlsx")
    parser.add_argument("--output-dir", default=".")
    args = parser.parse_args()

    workbook_path = Path(args.workbook)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(workbook_path, data_only=True, read_only=False)
    cards_ws = wb["Cards List"]
    drops_ws = wb["Drop Pools"]
    unlock_ws = wb["Drop Pool Unlock Requirements"]

    name_to_id, id_to_name = load_card_lookup(cards_ws)
    if len(name_to_id) != 722:
        raise ValueError(f"Expected 722 cards, found {len(name_to_id)}")

    sections = find_duelist_sections(drops_ws)
    if len(sections) != 39:
        raise ValueError(f"Expected 39 duelists, found {len(sections)}")

    duelists = [
        {"id": i + 1, "name": name, "slug": slugify(name)}
        for i, (_, name) in enumerate(sections)
    ]

    unlock_requirements = parse_unlock_requirements(unlock_ws)
    unlock_output = []
    drops: list[dict[str, Any]] = []
    pool_validation: list[dict[str, Any]] = []

    for index, (duelist_row, duelist_name) in enumerate(sections):
        duelist_id = index + 1
        header_row = duelist_row + 1
        end_row = (sections[index + 1][0] - 1) if index + 1 < len(sections) else drops_ws.max_row

        # Find the base and optional unlock pools from the header row.
        headers: dict[str, int] = {}
        for col in (1, 3, 5, 7):
            header = drops_ws.cell(header_row, col).value
            if header:
                headers[str(header).strip()] = col

        base_pools: dict[str, dict[int, dict[str, Any]]] = {}
        for base_header, rank in RANK_MAP.items():
            if base_header not in headers:
                raise ValueError(f"{duelist_name}: missing {base_header} pool")
            pool = read_pool(
                drops_ws, header_row + 1, end_row, headers[base_header], rank, name_to_id
            )
            total = sum(item["weight"] for item in pool.values())
            if abs(total - DROP_DENOMINATOR) > 1e-9:
                raise ValueError(f"{duelist_name} {base_header} totals {total}, expected {DROP_DENOMINATOR}")
            base_pools[rank] = pool

        # Start with every default/base pool row.
        base_drop_index: dict[tuple[str, int], int] = {}
        for rank, pool in base_pools.items():
            for card_id, item in pool.items():
                base_drop_index[(rank, card_id)] = len(drops)
                drops.append(make_drop(duelist_id, rank, item))

        # Handle alternate unlock pool, if one exists.
        unlock_headers = [(h, c) for h, c in headers.items() if h.endswith("(Unlock)")]
        if len(unlock_headers) > 1:
            raise ValueError(f"{duelist_name}: multiple unlock pools are not supported")

        if unlock_headers:
            unlock_header, unlock_col = unlock_headers[0]
            base_header = unlock_header.replace(" (Unlock)", "")
            if base_header not in RANK_MAP:
                raise ValueError(f"{duelist_name}: unknown unlock rank {unlock_header!r}")
            rank = RANK_MAP[base_header]
            unlock_pool = read_pool(drops_ws, header_row + 1, end_row, unlock_col, rank, name_to_id)
            total = sum(item["weight"] for item in unlock_pool.values())
            if abs(total - DROP_DENOMINATOR) > 1e-9:
                raise ValueError(f"{duelist_name} {unlock_header} totals {total}, expected {DROP_DENOMINATOR}")

            reqs = unlock_requirements.get(duelist_name)
            if not reqs:
                raise ValueError(f"{duelist_name}: unlock pool exists but requirements sheet has no entry")
            unlock_condition = "Unlock pool when: " + "; ".join(reqs)
            default_condition = "Default pool only; replaced when unlock requirements are met: " + "; ".join(reqs)
            changed_default_condition = "Default pool weight; changes when unlock requirements are met: " + "; ".join(reqs)

            base_pool = base_pools[rank]
            base_ids = set(base_pool)
            unlock_ids = set(unlock_pool)

            removed = sorted(base_ids - unlock_ids)
            added = sorted(unlock_ids - base_ids)
            changed = sorted(
                card_id for card_id in base_ids & unlock_ids
                if base_pool[card_id]["weight"] != unlock_pool[card_id]["weight"]
            )

            # Mark default rows that disappear or change under the unlock pool.
            for card_id in removed:
                pos = base_drop_index[(rank, card_id)]
                drops[pos]["condition"] = combine_conditions(
                    drops[pos].get("condition"), default_condition
                )
            for card_id in changed:
                pos = base_drop_index[(rank, card_id)]
                drops[pos]["condition"] = combine_conditions(
                    drops[pos].get("condition"), changed_default_condition
                )

            # Add only the unlock rows that differ from the base pool.
            for card_id in added + changed:
                drops.append(make_drop(duelist_id, rank, unlock_pool[card_id], unlock_condition))

            unlock_output.append({
                "duelistId": duelist_id,
                "duelistName": duelist_name,
                "rank": rank,
                "requirements": reqs,
                "added": [
                    {"cardId": c, "name": id_to_name[c], "weight": unlock_pool[c]["weight"]}
                    for c in added
                ],
                "removed": [
                    {"cardId": c, "name": id_to_name[c], "weight": base_pool[c]["weight"]}
                    for c in removed
                ],
                "changed": [
                    {
                        "cardId": c,
                        "name": id_to_name[c],
                        "defaultWeight": base_pool[c]["weight"],
                        "unlockWeight": unlock_pool[c]["weight"],
                    }
                    for c in changed
                ],
            })

            pool_validation.append({
                "duelist": duelist_name,
                "rank": rank,
                "defaultTotal": sum(x["weight"] for x in base_pool.values()),
                "unlockTotal": sum(x["weight"] for x in unlock_pool.values()),
            })

    # Stable ordering for predictable diffs.
    rank_order = {"SA_POW": 0, "BCD": 1, "SA_TEC": 2}
    drops.sort(key=lambda d: (d["duelistId"], rank_order[d["rank"]], d["cardId"], 1 if "Unlock pool when:" in d.get("condition", "") else 0))

    # Structural validation.
    if len({d["id"] for d in duelists}) != 39:
        raise ValueError("Duelist IDs are not unique")
    if len({d["slug"] for d in duelists}) != 39:
        raise ValueError("Duelist slugs are not unique")
    bad_cards = [d for d in drops if d["cardId"] not in id_to_name]
    if bad_cards:
        raise ValueError(f"Found drops referencing unknown card IDs: {bad_cards[:3]}")

    (output_dir / "duelists.json").write_text(
        json.dumps(duelists, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    (output_dir / "drops.json").write_text(
        json.dumps(drops, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    (output_dir / "drop-unlock-requirements.json").write_text(
        json.dumps(unlock_output, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )

    special_counts = defaultdict(int)
    for d in drops:
        condition = d.get("condition", "")
        if "400 wins" in condition:
            special_counts["400-win"] += 1
        if "already be in the library" in condition:
            special_counts["library-required"] += 1
        if "Kuriboh POW questline" in condition:
            special_counts["kuriboh-questline"] += 1
        if "Unlock pool when:" in condition:
            special_counts["unlock-variant"] += 1

    print(f"Extracted {len(duelists)} duelists")
    print(f"Extracted {len(drops)} drop records")
    print(f"Unlock pool variants: {len(unlock_output)} duelists")
    print(f"Special-condition counts: {dict(special_counts)}")
    for row in pool_validation:
        print(
            f"  {row['duelist']} {row['rank']}: "
            f"default={row['defaultTotal']}, unlock={row['unlockTotal']}"
        )
    print(f"Wrote {output_dir / 'duelists.json'}")
    print(f"Wrote {output_dir / 'drops.json'}")
    print(f"Wrote {output_dir / 'drop-unlock-requirements.json'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
